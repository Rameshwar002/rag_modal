"""
ingestion/loader.py
-------------------
Load raw text from files, Confluence, and SharePoint.
"""
from __future__ import annotations

import csv
import io
import json
import os
from dataclasses import dataclass, field
from pathlib import Path

from src.utils.logger import get_logger

log = get_logger(__name__)


@dataclass
class Document:
    text:     str
    source:   str        # "file" | "confluence" | "sharepoint"
    title:    str = ""
    metadata: dict = field(default_factory=dict)

    def __repr__(self):
        return f"Document(title={self.title!r}, chars={len(self.text)}, source={self.source!r})"


# ── File loaders ──────────────────────────────────────────────────────────────

def _load_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)

def _load_docx(path: Path) -> str:
    from docx import Document as DocxDoc
    doc = DocxDoc(str(path))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())

def _load_csv(path: Path) -> str:
    rows = []
    with open(path, newline="", encoding="utf-8", errors="ignore") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(" | ".join(f"{k}: {v}" for k, v in row.items()))
    return "\n".join(rows)

def _load_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8"))
    return json.dumps(data, indent=2)

def _load_html(path: Path) -> str:
    from bs4 import BeautifulSoup
    return BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser").get_text(separator="\n")

LOADERS = {
    ".txt": _load_txt, ".md":   _load_txt, ".yaml": _load_txt,
    ".yml": _load_txt, ".log":  _load_txt,
    ".pdf": _load_pdf,
    ".docx":_load_docx,
    ".csv": _load_csv,
    ".json":_load_json,
    ".html":_load_html, ".htm": _load_html,
}


def load_file(path: str | Path) -> Document:
    path = Path(path)
    ext  = path.suffix.lower()
    if ext not in LOADERS:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {list(LOADERS)}")
    text = LOADERS[ext](path)
    log.info(f"file_loaded  path={path}  chars={len(text)}")
    return Document(
        text=text, source="file", title=path.name,
        metadata={"file_path": str(path), "extension": ext},
    )


def load_files_from_dir(directory: str | Path, recursive: bool = True) -> list[Document]:
    directory = Path(directory)
    pattern   = "**/*" if recursive else "*"
    docs = []
    for p in directory.glob(pattern):
        if p.is_file() and p.suffix.lower() in LOADERS:
            try:
                docs.append(load_file(p))
            except Exception as e:
                log.warning(f"file_load_failed  path={p}  error={e}")
    return docs


# ── Confluence ────────────────────────────────────────────────────────────────

def load_confluence(
    space_key: str,
    url:       str | None = None,
    username:  str | None = None,
    api_token: str | None = None,
    limit:     int = 200,
) -> list[Document]:
    from atlassian import Confluence
    from bs4 import BeautifulSoup

    url       = url       or os.environ["CONFLUENCE_URL"]
    username  = username  or os.environ["CONFLUENCE_USERNAME"]
    api_token = api_token or os.environ["CONFLUENCE_API_TOKEN"]

    confluence = Confluence(url=url, username=username, password=api_token)
    pages      = confluence.get_all_pages_from_space(
        space_key, start=0, limit=limit, expand="body.storage"
    )
    docs = []
    for page in pages:
        try:
            html = page["body"]["storage"]["value"]
            text = BeautifulSoup(html, "html.parser").get_text(separator="\n")
        except Exception:
            text = page.get("title", "")
        docs.append(Document(
            text=text, source="confluence", title=page.get("title", ""),
            metadata={
                "page_id": page.get("id"),
                "space":   space_key,
                "url":     f"{url}/wiki/spaces/{space_key}/pages/{page.get('id')}",
            },
        ))
    log.info(f"confluence_loaded  space={space_key}  pages={len(docs)}")
    return docs


# ── SharePoint ────────────────────────────────────────────────────────────────

def load_sharepoint(
    site_url:      str | None = None,
    folder_path:   str        = "/Shared Documents",
    client_id:     str | None = None,
    client_secret: str | None = None,
    tenant_id:     str | None = None,
) -> list[Document]:
    from office365.runtime.auth.client_credential import ClientCredential
    from office365.sharepoint.client_context import ClientContext

    site_url      = site_url      or os.environ["SHAREPOINT_URL"]
    client_id     = client_id     or os.environ["SHAREPOINT_CLIENT_ID"]
    client_secret = client_secret or os.environ["SHAREPOINT_CLIENT_SECRET"]

    ctx    = ClientContext(site_url).with_credentials(ClientCredential(client_id, client_secret))
    folder = ctx.web.get_folder_by_server_relative_url(folder_path)
    files  = folder.files
    ctx.load(files)
    ctx.execute_query()

    docs = []
    for f in files:
        try:
            response = f.read()
            ctx.execute_query()
            content  = response.value
            ext      = Path(f.properties["Name"]).suffix.lower()

            if ext == ".pdf":
                from pypdf import PdfReader
                text = "\n\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(content)).pages)
            elif ext == ".docx":
                from docx import Document as DocxDoc
                doc  = DocxDoc(io.BytesIO(content))
                text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            else:
                text = content.decode("utf-8", errors="ignore")

            docs.append(Document(
                text=text, source="sharepoint",
                title=f.properties["Name"],
                metadata={
                    "file_name":  f.properties["Name"],
                    "site_url":   site_url,
                    "folder":     folder_path,
                },
            ))
        except Exception as e:
            log.warning(f"sharepoint_file_failed  file={f.properties.get('Name')}  error={e}")

    log.info(f"sharepoint_loaded  folder={folder_path}  files={len(docs)}")
    return docs
